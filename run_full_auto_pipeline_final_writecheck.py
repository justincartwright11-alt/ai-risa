# Derivative of the final pipeline for write-path debugging
# Adds prints to trace section population and write conditions

import argparse
import os
import sys
import json
from docx import Document
from run_full_auto_pipeline_recovered_embed import resolve_bout_identity, is_true_tba_bout, build_fight_report

def debug_print(msg, *args):
    print(f"[WRITECHK] {msg}", *args)

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--mode', required=True, choices=['fighter', 'matchup', 'event'])
    parser.add_argument('--fighter', help='Fighter name for single dossier')
    parser.add_argument('--fighter-a', help='Fighter A for matchup')
    parser.add_argument('--fighter-b', help='Fighter B for matchup')
    parser.add_argument('--event', help='Event slug for event mode (e.g., boxing_wollongong_2026_04_05)')
    parser.add_argument('--report', help='Report type (Full Prim Report, Fighter/Trainer, Promoter, Broadcast, Premium Fan)')
    parser.add_argument('--format', help='Export format (DOCX or PDF)')
    args = parser.parse_args()
    mode = str(args.mode or '').strip().lower()
    report = str(args.report or '').strip().lower().replace('_', ' ')
    fmt = str(args.format or '').strip().lower()
    debug_print('Dispatch args', {'mode': mode, 'report': report, 'format': fmt})

    report_names = {'full prim report', 'full prim', 'full_prim'}
    if mode == 'event' and report.replace('_', ' ') in report_names:
        event_summary = None
        event_summary_path = f"C:/ai_risa_data/events/{args.event}.json"
        debug_print('Event summary path', event_summary_path)
        if os.path.exists(event_summary_path):
            with open(event_summary_path, 'r', encoding='utf-8') as f:
                event_summary = json.load(f)
        else:
            debug_print('Event summary file does not exist!')
        event_summary = event_summary or {}
        debug_print('Event summary keys', list(event_summary.keys()))
        bout_entries = (
            event_summary.get('results')
            or event_summary.get('fights')
            or event_summary.get('bouts')
            or event_summary.get('card')
            or []
        )
        if isinstance(bout_entries, dict):
            bout_entries = list(bout_entries.values())
        debug_print('Bout entries count', len(bout_entries))
        tba_sections = []
        premium_sections = []
        for i, bout in enumerate(bout_entries):
            fighters = bout.get('fighters') or []
            left = fighters[0] if len(fighters) > 0 else ''
            right = fighters[1] if len(fighters) > 1 else ''
            fight_name = f"{left} vs. {right}" if left and right else left or right or 'Unknown Bout'
            pred_path = bout.get('prediction_file')
            pred = {}
            if pred_path and os.path.exists(pred_path):
                with open(pred_path, 'r', encoding='utf-8') as pf:
                    pred = json.load(pf)
            fighter_a_profile = None
            fighter_b_profile = None
            fighter_a_profile_path = f"C:/ai_risa_data/fighters/{left.lower().replace(' ', '_')}.json"
            fighter_b_profile_path = f"C:/ai_risa_data/fighters/{right.lower().replace(' ', '_')}.json"
            if os.path.exists(fighter_a_profile_path):
                with open(fighter_a_profile_path, 'r', encoding='utf-8') as f:
                    fighter_a_profile = json.load(f)
            if os.path.exists(fighter_b_profile_path):
                with open(fighter_b_profile_path, 'r', encoding='utf-8') as f:
                    fighter_b_profile = json.load(f)
            matchup_data = bout
            fighter_a_name, fighter_b_name, bout_title = resolve_bout_identity(matchup_data, fight_name, pred, fighter_a_profile, fighter_b_profile)
            if is_true_tba_bout(fighter_a_name, fighter_b_name):
                tba_sections.append(bout_title)
                continue
            fight_report = build_fight_report(
                prediction=pred,
                fighter_a_profile=fighter_a_profile,
                fighter_b_profile=fighter_b_profile,
                matchup_data=matchup_data,
                event_summary=event_summary,
                event_fight_name=fight_name,
            )
            if fight_report:
                premium_sections.append(fight_report['bout_title'])
        debug_print('len(premium_sections)', len(premium_sections))
        debug_print('len(tba_sections)', len(tba_sections))
        md_path = f"C:/ai_risa_data/reports/{args.event}_full_prim_report.md"
        docx_path = f"C:/ai_risa_data/reports/{args.event}_full_prim_report.docx"
        pdf_path = f"C:/ai_risa_data/reports/{args.event}_full_prim_report.pdf"
        debug_print('Output paths', {'md': md_path, 'docx': docx_path, 'pdf': pdf_path})
        if premium_sections or tba_sections:
            debug_print('About to write markdown', md_path)
            with open(md_path, 'w', encoding='utf-8') as f:
                f.write('WRITECHK: dummy content')
            debug_print('About to write docx', docx_path)
            doc = Document()
            doc.add_paragraph('WRITECHK: dummy content')
            doc.save(docx_path)
            if fmt == 'pdf':
                debug_print('About to write pdf', pdf_path)
                from fpdf import FPDF
                pdf = FPDF()
                pdf.add_page()
                pdf.set_font('Arial', size=12)
                pdf.cell(200, 10, txt='WRITECHK: dummy content', ln=True)
                pdf.output(pdf_path)
        else:
            debug_print('No sections to write, skipping output')
    else:
        debug_print(f'Unknown mode or report type. mode={mode!r}, report={report!r}, format={fmt!r}')
        sys.exit(1)

if __name__ == "__main__":
    main()
